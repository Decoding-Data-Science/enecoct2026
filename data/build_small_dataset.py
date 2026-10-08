"""
Build the small Colab-friendly A-001 / Asset 360 dataset from the full
Nuclear Enterprise 360 V2.2 SQLite database (release asset, ~120 MB).

Usage:
    python build_small_dataset.py path/to/nuclear_enterprise_360_v2_2_clean.db

Output:
    data/csv/*.csv                      small tables loaded by the Colab notebooks
    data/excel/asset_360.xlsx           Copilot-ready workbook (asset view + A-001 sheets)
    data/excel/project_360.xlsx         Copilot-ready workbook (project view)

The CSV files are already committed, so participants never need the big database.
Only run this script if you want to regenerate the small dataset.
"""
import sqlite3, sys, os
from datetime import datetime
import pandas as pd

DB = sys.argv[1] if len(sys.argv) > 1 else "nuclear_enterprise_360_v2_2_clean.db"
HERE = os.path.dirname(os.path.abspath(__file__))
CSV = os.path.join(HERE, "csv")
XLS = os.path.join(HERE, "excel")
os.makedirs(CSV, exist_ok=True)
os.makedirs(XLS, exist_ok=True)
con = sqlite3.connect(DB)
q = lambda sql: pd.read_sql(sql, con)

HOURLY_DAYS = 90   # last 90 days of hourly A-001 data = 2,160 rows

# ---------------------------------------------------------------- small CSV set
hourly = q("select * from a001_sensor_hourly order by hour_timestamp")
hourly["hour_timestamp"] = pd.to_datetime(hourly["hour_timestamp"])
cut = hourly["hour_timestamp"].max() - pd.Timedelta(days=HOURLY_DAYS) + pd.Timedelta(hours=1)
hourly_small = hourly[hourly["hour_timestamp"] >= cut].copy()

tables = {
    "assets":                  q("select * from assets order by asset_id"),
    "systems":                 q("select * from systems order by system_id"),
    "asset_360":               q("select * from asset_360 order by asset_id"),
    "a001_sensor_hourly_90d":  hourly_small,
    "a001_event_timeline":     q("select * from a001_event_timeline order by event_date, event_id"),
    "a001_asset_health_scores":q("select * from asset_health_scores where asset_id='A-001' order by score_date"),
    "work_orders":             q("select * from work_orders order by work_order_id"),
    "inspections":             q("select * from inspections order by inspection_id"),
    "maintenance_history":     q("select * from maintenance_history order by maintenance_id"),
    "a001_project_risk_360":   q("select * from a001_project_risk_360 order by project_id, risk_id"),
    "a001_source_documents":   q("select * from a001_source_documents order by document_id"),
    "a001_document_evidence":  q("select * from a001_document_evidence order by document_id"),
}
for name, df in tables.items():
    df.to_csv(os.path.join(CSV, name + ".csv"), index=False)
    print(f"csv  {name:28s} {len(df):>6,d} rows")

# ---------------------------------------------------------------- Excel workbooks (Copilot-ready)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

HDR = PatternFill("solid", fgColor="1F3864")

def add_sheet(wb, name, df, widths=None, note=None):
    ws = wb.create_sheet(name)
    df = df.copy()
    for c in df.columns:
        if c.endswith("_date") or c in ("hour_timestamp", "event_date", "score_date", "effective_date", "evidence_date", "kpi_month"):
            parsed = pd.to_datetime(df[c], errors="coerce")
            if parsed.notna().sum() > 0:
                df[c] = parsed
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        vals = []
        for v in row:
            if pd.isna(v):
                vals.append(None)
            elif isinstance(v, pd.Timestamp):
                vals.append(v.to_pydatetime())
            else:
                vals.append(v.item() if hasattr(v, "item") else v)
        ws.append(vals)
    n_rows, n_cols = len(df) + 1, len(df.columns)
    tname = "".join(ch if ch.isalnum() or ch == "_" else "_" for ch in name)
    t = Table(displayName="t_" + tname, ref=f"A1:{get_column_letter(n_cols)}{max(n_rows, 2)}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HDR
        letter = get_column_letter(j)
        longest = max([len(str(c))] + [len(str(x)) for x in df[c].head(200).tolist()])
        ws.column_dimensions[letter].width = min(max(10, longest + 2), 60)
        if str(df[c].dtype).startswith("datetime"):
            fmt = "yyyy-mm-dd hh:mm" if c == "hour_timestamp" else ("yyyy-mm" if c == "kpi_month" else "yyyy-mm-dd")
            for r in range(2, n_rows + 1):
                ws.cell(row=r, column=j).number_format = fmt
    ws.freeze_panes = "A2"
    return ws

def readme_sheet(wb, title, intro, sheets, prompts, notes):
    ws = wb.active; ws.title = "README"
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=16, color="1F3864")
    r = 3
    for line in intro:
        ws.cell(row=r, column=1, value=line); r += 1
    r += 1
    ws.cell(row=r, column=1, value="Sheet").font = Font(bold=True)
    ws.cell(row=r, column=2, value="Rows").font = Font(bold=True)
    ws.cell(row=r, column=3, value="What it contains").font = Font(bold=True)
    r += 1
    for s, n, d in sheets:
        ws.cell(row=r, column=1, value=s); ws.cell(row=r, column=2, value=n); ws.cell(row=r, column=3, value=d); r += 1
    r += 1
    ws.cell(row=r, column=1, value="Suggested prompts for Copilot / ChatGPT / Colab").font = Font(bold=True, size=12); r += 1
    for i, p in enumerate(prompts, 1):
        ws.cell(row=r, column=1, value=f"{i}."); ws.cell(row=r, column=2, value=p); r += 1
    r += 1
    for n in notes:
        ws.cell(row=r, column=1, value="Note"); ws.cell(row=r, column=2, value=n); r += 1
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 12; ws.column_dimensions["C"].width = 110

def dictionary_sheet(wb, rows):
    df = pd.DataFrame(rows, columns=["sheet", "column", "meaning"])
    add_sheet(wb, "data_dictionary", df)

# ---- workbook 1: asset_360
wb = Workbook()
a360 = tables["asset_360"]
apm = q("select * from asset_project_map order by project_id, asset_id")
aprm = q("select * from asset_project_risk_map order by asset_id, project_id, risk_id")
sheets1 = [
    ("asset_360", a360, "MAIN VIEW: one row per asset (128) with latest health score, risk level, open and high-priority work orders, follow-up findings."),
    ("asset_project_map", apm, "Which assets sit in which projects."),
    ("asset_project_risk_map", aprm, "Project and direct asset risks per asset."),
    ("a001_project_risk_360", tables["a001_project_risk_360"], "A-001 only: its projects and their risks with evidence basis."),
    ("a001_event_timeline", tables["a001_event_timeline"], "A-001 only: chronological events (project start, work orders, inspections, maintenance, documents)."),
    ("a001_document_evidence", tables["a001_document_evidence"], "A-001 source documents (procedures, versions, approval status) and what each links to. Includes SUPERSEDED and DRAFT versions."),
    ("a001_sensor_hourly", hourly, "A-001 sensor readings aggregated to hourly from 525,600 one-minute readings (full year, 8,760 rows)."),
]
readme_sheet(
    wb, "ENEC 2026 - Asset 360 workbook (synthetic training data)",
    ["Synthetic data from Nuclear Enterprise 360 V2.2. Focus asset: A-001 (cooling-water pump).",
     "Every sheet is an Excel Table with a single header row, real Excel dates and static values, so Microsoft Copilot in Excel can analyse it directly."],
    [(n, len(d), desc) for n, d, desc in sheets1] + [("data_dictionary", "", "Column descriptions for the main sheets.")],
    ["Which assets are HIGH criticality with a health score below 70, and how many open or high-priority work orders does each have?",
     "Summarise asset_360 by system_name and unit_id: average health score, count by risk_level, total open work orders. Show as a chart.",
     "Which assets have follow-up findings but no high-priority open work? Flag them as potential gaps.",
     "For A-001, build a narrative from a001_event_timeline: what happened, in what order, and what is still open?",
     "Using a001_sensor_hourly, find when A-001 vibration started to rise and how it relates to anomaly_minutes. Chart it.",
     "In a001_document_evidence, which documents are SUPERSEDED or DRAFT, and what is the risk of an AI agent citing them?",
     "Join asset_project_risk_map with asset_360: which assets sit on projects with the highest-exposure risks?"],
    ["asset_360 is a database view; the supporting sheets are also database views, extracted as-is. Date columns are real Excel dates.",
     "Training data only. Nothing here is an engineering diagnosis or a maintenance authorisation."])
for n, d, _ in sheets1:
    add_sheet(wb, n, d)
D = [("asset_360", "asset_id", "Unique asset identifier (A-001 is the focus asset: cooling-water pump)"),
     ("asset_360", "criticality", "LOW / MEDIUM / HIGH"),
     ("asset_360", "system_name", "Plant system the asset belongs to"),
     ("asset_360", "unit_id", "Training unit (e.g. TRN-A)"),
     ("asset_360", "health_score", "Latest model health score, 0-100 (higher is healthier)"),
     ("asset_360", "risk_level", "Risk level from the latest health scoring"),
     ("asset_360", "recommended_action", "Action recommended by the latest health scoring"),
     ("asset_360", "open_work_orders", "Work orders with status OPEN, IN_PROGRESS or DEFERRED"),
     ("asset_360", "high_priority_open_work", "Open work orders with priority HIGH or URGENT"),
     ("asset_360", "follow_up_findings", "Inspections flagged follow_up_required"),
     ("asset_360", "latest_inspection_date", "Date of the latest follow-up-required inspection (not necessarily the latest inspection)"),
     ("asset_project_risk_map", "exposure_score", "probability x impact (each 1-5)"),
     ("asset_project_risk_map", "risk_scope", "DIRECT_TO_ASSET if a risk_assets record ties the risk to this asset, else PROJECT_LEVEL"),
     ("a001_event_timeline", "event_type", "PROJECT_START, WORK_ORDER_CREATED, WORK_ORDER_COMPLETED, INSPECTION, MAINTENANCE, DOCUMENT"),
     ("a001_document_evidence", "approval_status", "APPROVED, SUPERSEDED, DRAFT or OPEN. Only APPROVED is current authority."),
     ("a001_sensor_hourly", "anomaly_minutes", "Minutes in the hour flagged anomalous"),
     ("a001_sensor_hourly", "quality_issue_minutes", "Minutes in the hour where sensor_quality was not GOOD")]
dictionary_sheet(wb, D)
wb.active.sheet_properties.tabColor = "1F3864"
wb.save(os.path.join(XLS, "asset_360.xlsx"))
print("xlsx asset_360.xlsx")

# ---- workbook 2: project_360
wb = Workbook()
p360 = q("""select p.project_id, p.project_name, d.department_name, p.project_manager_code, p.strategic_theme,
  p.start_date, p.planned_end_date, p.status as project_status, p.completion_pct, p.budget_usd,
  v.active_risks, v.maximum_exposure, v.active_actions, v.overdue_actions,
  (select count(*) from project_assets pa where pa.project_id=p.project_id) as linked_assets,
  (select count(*) from risks r where r.project_id=p.project_id) as total_risks
  from projects p join departments d on d.department_id=p.department_id
  join project_risk_360 v on v.project_id=p.project_id order by p.project_id""")
kpis = q("select k.*, p.project_name from project_kpis k join projects p on p.project_id=k.project_id order by k.project_id, k.kpi_month")
risks = q("select r.*, p.project_name from risks r join projects p on p.project_id=r.project_id order by r.risk_id")
acts = q("select a.*, p.project_name from actions a join projects p on p.project_id=a.project_id order by a.action_id")
pas = q("select pa.*, p.project_name, s.asset_name from project_assets pa join projects p on p.project_id=pa.project_id join assets s on s.asset_id=pa.asset_id order by pa.project_id")
sheets2 = [
    ("project_360", p360, "MAIN VIEW: one row per project (29) with department, manager code, theme, dates, budget, completion, risk and action counts."),
    ("project_kpis", kpis, "Monthly planned vs actual progress and cost per project (includes future-dated months)."),
    ("risks", risks, "All project risks with probability, impact, exposure and mitigation plan."),
    ("actions", acts, "Mitigation actions with owner, priority, due date and status."),
    ("project_asset_links", pas, "Which assets each project covers (only a few projects have linked assets)."),
]
readme_sheet(
    wb, "ENEC 2026 - Project 360 workbook (synthetic training data)",
    ["Synthetic data from Nuclear Enterprise 360 V2.2.",
     "project_360 is built from the database view project_risk_360 joined with project attributes."],
    [(n, len(d), desc) for n, d, desc in sheets2],
    ["Which projects are over budget or behind plan, and which have overdue actions?",
     "Rank projects by maximum_exposure and show the top 5 with their risk categories.",
     "Which risks have no mitigation action or an overdue action?",
     "Chart planned vs actual progress by month for the five largest-budget projects.",
     "Which projects are linked to assets, and what do their risks say about those assets?"],
    ["Only P-021, P-027 and P-029 link to assets. Questions that join projects to assets will mostly return empty results for other projects.",
     "project_kpis includes future-dated months. Training data only."])
for n, d, _ in sheets2:
    add_sheet(wb, n, d)
wb.active.sheet_properties.tabColor = "1F3864"
wb.save(os.path.join(XLS, "project_360.xlsx"))
print("xlsx project_360.xlsx")
